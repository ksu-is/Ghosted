from datetime import datetime
from openpyxl import load_workbook

from db import get_db_connection

EXPECTED_COLUMNS = {
    "company": "company",
    "position": "role",
    "role": "role",
    "location": "location",
    "date_applied": "date_applied",
    "status": "status",
    "interview_date": "interview_date",
    "follow_up_date": "follow_up_date",
    "offer_status": "offer_status",
    "notes": "notes",
}


def import_excel_file(file_path: str) -> tuple[int, int]:
    workbook = load_workbook(filename=file_path)
    sheet = workbook.active

    header_row = [clean_header(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    column_map = map_columns(header_row)

    imported_count = 0
    skipped_rows = 0

    conn = get_db_connection()

    for row in sheet.iter_rows(min_row=2, values_only=True):
        row_data = build_row_dict(row, column_map)
        if not row_data.get("company") or not row_data.get("role"):
            skipped_rows += 1
            continue

        conn.execute(
            """
            INSERT INTO applications
            (company, role, location, date_applied, status, interview_date, follow_up_date, offer_status, notes)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                row_data.get("company", ""),
                row_data.get("role", ""),
                row_data.get("location", ""),
                normalize_date(row_data.get("date_applied")),
                row_data.get("status", "Applied") or "Applied",
                normalize_date(row_data.get("interview_date")),
                normalize_date(row_data.get("follow_up_date")),
                row_data.get("offer_status", "Pending") or "Pending",
                row_data.get("notes", ""),
            ),
        )
        imported_count += 1

    conn.commit()
    conn.close()
    workbook.close()
    return imported_count, skipped_rows


def clean_header(value) -> str:
    if value is None:
        return ""
    return str(value).strip().lower().replace(" ", "_")


def map_columns(headers: list[str]) -> dict[int, str]:
    result = {}
    for index, header in enumerate(headers):
        if header in EXPECTED_COLUMNS:
            result[index] = EXPECTED_COLUMNS[header]
    return result


def build_row_dict(row: tuple, column_map: dict[int, str]) -> dict[str, str]:
    result = {}
    for index, field_name in column_map.items():
        value = row[index] if index < len(row) else ""
        result[field_name] = "" if value is None else str(value).strip()
    return result


def normalize_date(value) -> str:
    if not value:
        return ""

    text = str(value).strip()
    for pattern in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y"):
        try:
            return datetime.strptime(text, pattern).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return text
